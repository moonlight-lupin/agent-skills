"""Contract tests for issue #13: fill-template extract direction (document → template).

Written by the orchestrator (Hermes) — builders must NOT modify this file.

Round-trip guarantee: a filled/generated document, run through extract_template,
yields a tokenised template + a data-table skeleton; filling that template with
the extracted rows reproduces the original values. Also: confirmation digest —
extract reports a mapping that the user must be able to review before reuse.

Scope this round: .docx only.
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import pytest
from docx import Document

REPO_ROOT = Path(__file__).resolve().parent.parent
ENGINE_DIR = REPO_ROOT / "productivity" / "fill-template" / "scripts"
sys.path.insert(0, str(ENGINE_DIR))

import fill_template as ft  # noqa: E402


# ─── helpers ─────────────────────────────────────────────────────────────────

def make_filled_docx(path: Path) -> None:
    """A realistic filled letter: boilerplate + per-recipient values."""
    doc = Document()
    doc.add_paragraph("Dear Alice Tan,")
    body = doc.add_paragraph(
        "Thank you for your payment of $420.00 on 12 Mar 2026 for invoice "
        "INV-1024. Your account ACME-0042 is now settled."
    )
    doc.add_paragraph("Dear Bob Lim,")
    doc.add_paragraph(
        "Thank you for your payment of $180.50 on 03 Apr 2026 for invoice "
        "INV-1025. Your account ACME-0043 is now settled."
    )
    doc.save(str(path))


# ─── the extract contract ────────────────────────────────────────────────────

def test_extract_template_exists_and_runs():
    """fill_template must expose extract_template(src, out_dir, **kw)."""
    assert hasattr(ft, "extract_template"), (
        "extract_template is the issue-#13 contract entry point; missing entirely"
    )


def test_extract_produces_template_and_skeleton(tmp_path):
    """extract → template .docx + data skeleton (.csv/.xlsx) exist and parse."""
    src = tmp_path / "filled.docx"
    make_filled_docx(src)
    out = tmp_path / "extracted"
    report = ft.extract_template(str(src), str(out))
    assert isinstance(report, dict), "extract_template must return a report dict"
    files = list(out.iterdir())
    templates = [f for f in files if f.suffix == ".docx"]
    tables = [f for f in files if f.suffix in (".csv", ".xlsx")]
    assert templates, "no tokenised template .docx produced"
    assert tables, "no data-table skeleton produced"
    # template contains tokens
    tmpl_text = ft.read_content(str(templates[0]))
    assert "{{" in tmpl_text and "}}" in tmpl_text, \
        "template must contain {{tokens}}"


def test_extract_tokens_match_existing_token_syntax():
    """Tokens emitted must parse with the engine's own TOKEN_RE — same syntax."""
    src = tmp_path = None  # placeholder to keep naming clear
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        src = tdp / "filled.docx"
        make_filled_docx(src)
        out = tdp / "extracted"
        ft.extract_template(str(src), str(out))
        tmpl = next(p for p in out.iterdir() if p.suffix == ".docx")
        tokens = ft.tokens_in(str(tmpl))
        assert tokens, "no tokens in extracted template"
        # engine token names: no spaces, braces closed
        for t in tokens:
            assert ft.TOKEN_RE.fullmatch("{{" + t + "}}" ), f"bad token {t!r}"


def test_round_trip_extract_then_fill(tmp_path):
    """THE guarantee: extract → fill with the skeleton values reproduces them."""
    src = tmp_path / "filled.docx"
    make_filled_docx(src)
    out = tmp_path / "extracted"
    ft.extract_template(str(src), str(out))

    tmpl = next(p for p in out.iterdir() if p.suffix == ".docx")
    table = next(p for p in out.iterdir() if p.suffix in (".csv", ".xlsx"))
    rows = ft.load_rows(str(table))
    assert rows, "skeleton has no rows"

    filled_dir = tmp_path / "refilled"
    report = ft.generate(str(tmpl), rows, token_to_column={}, outdir=str(filled_dir))
    assert report, "generate returned nothing"

    # every per-row value must reappear in the refilled documents
    combined = ""
    for f in sorted(filled_dir.iterdir()):
        combined += ft.read_content(str(f))
    for needle in ["Alice Tan", "Bob Lim", "$420.00", "$180.50",
                   "12 Mar 2026", "03 Apr 2026", "INV-1024", "INV-1025",
                   "ACME-0042", "ACME-0043"]:
        assert needle in combined, f"round trip lost {needle!r}"


def test_extract_report_lists_mapping(tmp_path):
    """Report must expose the inferred mapping for user confirmation (the gate)."""
    src = tmp_path / "filled.docx"
    make_filled_docx(src)
    out = tmp_path / "extracted"
    report = ft.extract_template(str(src), str(out))
    # the mapping must be in the report or alongside it as a file the agent can show
    has_mapping = (
        ("mapping" in report and report["mapping"])
        or any("mapping" in f.name.lower() for f in out.iterdir())
    )
    assert has_mapping, "extract must surface the inferred token mapping"


def test_extract_never_edits_source(tmp_path):
    """Workspace hygiene: the source document must be byte-identical after extract."""
    src = tmp_path / "filled.docx"
    make_filled_docx(src)
    before = src.read_bytes()
    ft.extract_template(str(src), str(tmp_path / "extracted"))
    assert src.read_bytes() == before, "extract_template modified the source file"


def test_extract_handles_docx_without_repeats(tmp_path):
    """A single-instance letter must still extract (one skeleton row), not crash."""
    doc = Document()
    doc.add_paragraph("Invoice total: $99.00 for ACME-0001.")
    src = tmp_path / "single.docx"
    doc.save(src)
    out = tmp_path / "extracted"
    report = ft.extract_template(str(src), str(out))
    assert isinstance(report, dict)
    assert list(out.iterdir()), "no outputs for single-instance document"